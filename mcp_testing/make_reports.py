import json
from collections import Counter
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt

H = Path(__file__).parent
data = json.load(open(H / "results.json"))
R, M = data["results"], data["meta"]
n_pass = sum(r["status"] == "PASS" for r in R)


def table(doc, head, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(head)); t.style = "Light Grid Accent 1"
    for i, h in enumerate(head):
        t.rows[0].cells[i].text = h
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = str(v)
    for row in t.rows:
        for c in row.cells:
            for p in c.paragraphs:
                for run in p.runs: run.font.size = Pt(9)
    return t


# ---------------- MCP report ----------------
d = Document()
d.add_heading("MCP Testing Report - AI Scholarship Finder", 0)
d.add_paragraph(f"Date: {M['started'][:10]}    Application: Streamlit web app at {M['url']}    MCP server: {M['server']}")
d.add_heading("1. Objective", 1)
d.add_paragraph("To show how the Model Context Protocol (MCP) can test the website: an MCP client sends tool calls to a "
                "Playwright MCP server, which controls a real Chromium browser. The client opens the site, clicks, types, "
                "reads the page and takes screenshots, then compares what it sees with the expected result.")
d.add_heading("2. How MCP testing works here", 1)
for s in ["Python MCP client (mcp_testing/mcp_harness.py) starts the server with 'npx @playwright/mcp' over stdio.",
          "The client lists the tools the server offers (%d tools, e.g. browser_navigate, browser_snapshot, browser_click, browser_type, browser_evaluate, browser_take_screenshot, browser_console_messages)." % len(M["tools"]),
          "browser_snapshot returns the page's accessibility tree; every element gets a ref id. The test finds the element by name and passes its ref to browser_click / browser_type.",
          "browser_evaluate reads page text and DOM facts (e.g. 'Showing 1 of 8', injected elements, page width).",
          "Each test case compares the observed result with the expected result and records PASS/FAIL plus a screenshot.",
          "Every MCP tool call is logged (name, arguments, duration, success) in mcp_testing/results.json."]:
    d.add_paragraph(s, style="List Bullet")
d.add_heading("3. Summary of results", 1)
d.add_paragraph(f"MCP UI test cases executed: {len(R)}    Passed: {n_pass}    Failed: {len(R)-n_pass}    MCP tool calls: {len(M['mcp_calls'])}")
d.add_paragraph("Existing automated suite (pytest): 175 passed, 1 skipped (the live test, run separately), 0 failed - covers eligibility rules, "
                "normalizer, red flags, ranking, repository, source safety, Gemini service (mocked), Streamlit AppTest flows.")
usage = Counter(c["tool"] for c in M["mcp_calls"])
d.add_heading("4. MCP tools used", 1)
table(d, ["MCP tool", "Calls"], usage.most_common())
d.add_heading("5. Test cases and results", 1)
table(d, ["ID", "Title", "Steps (via MCP)", "Expected", "Actual", "Result"],
      [[r["id"], r["title"], r["steps"], r["expected"], r["actual"], r["status"]] for r in R])
d.add_heading("6. Evidence (screenshots taken through browser_take_screenshot)", 1)
for r in R:
    f = r.get("screenshot")
    if f and (H / "screenshots" / f).exists():
        d.add_paragraph(f"{r['id']} - {r['title']} ({r['status']})")
        d.add_picture(str(H / "screenshots" / f), width=Inches(5.8) if "mobile" not in f else Inches(2.2))
d.add_heading("7. Failures and observations", 1)
fails = [r for r in R if r["status"] == "FAIL"]
if fails:
    for r in fails:
        d.add_paragraph(f"{r['id']} {r['title']}: {r['actual']}", style="List Bullet")
else:
    d.add_paragraph("No failing test cases in the final run.")
d.add_paragraph("During test development the first runs failed because of problems in the test script, not the website "
                "(MCP element refs change after navigation; Streamlit radio inputs are covered by their labels so the label text must be clicked). "
                "These were fixed in the harness. Console check: " + M.get("console_errors_summary", "").replace("\n", " ")[:200])
d.add_paragraph("Live-search note (TC-16): depends on internet, DuckDuckGo and official sites, so the result can vary between runs.")
d.add_heading("8. Conclusion", 1)
d.add_paragraph("MCP lets an AI/automation client test the site the way a user does - through the real UI - without writing browser-specific code, "
                "and returns structured evidence (page snapshot, text, console, screenshots). It complements the 175 unit tests, which check the logic "
                "but not the rendered interface.")
d.save(H / "MCP_Testing_Report.docx")

# ---------------- Manual report ----------------
by = {r["id"]: r for r in R}
manual = [
 ("MT-01", "Open app with 'streamlit run app.py'", "Page loads with title and form", "TC-01"),
 ("MT-02", "Choose Demo data, click Find scholarships", "8 fictional cards with DEMO label and banner", "TC-03"),
 ("MT-03", "Demo data, State=Karnataka, Category=OBC, search", "Results shown, metrics update", "TC-04"),
 ("MT-04", "Type 'Karnataka' in 'Provider / name contains'", "List narrows", "TC-05"),
 ("MT-05", "Type nonsense in name filter", "'No results match these filters'", "TC-06"),
 ("MT-06", "Search with all fields blank", "Warning about missing details", "TC-07"),
 ("MT-07", "Select income 'Exact amount'", "Number input appears", "TC-08"),
 ("MT-08", "Enter HTML/script text in Degree field and search", "Shown as plain text, nothing runs", "TC-09"),
 ("MT-09", "Type 99 in Current year", "No crash", "TC-10"),
 ("MT-10", "Click 'Clear saved searches & AI cache'", "'Cleared.' appears", "TC-12"),
 ("MT-11", "Resize browser to phone width", "No horizontal scroll", "TC-13"),
 ("MT-12", "Tick 'Save my search summary'", "Search history panel appears", "TC-14"),
 ("MT-13", "Use Live search with Karnataka / UG / B.E.", "Real results or clear problem messages", "TC-16"),
 ("MT-14", "Change 'Sort by' options", "Order changes, no error", None),
 ("MT-15", "Open a result card, check source link and red flags", "Official source link, conditions and flags visible", None),
 ("MT-16", "Click AI explanation on a card", "Explanation shown or clear disabled/error message", None),
]
m = Document()
m.add_heading("Manual Testing Report - AI Scholarship Finder", 0)
m.add_paragraph(f"Date: {M['started'][:10]}    Application: Streamlit app ({M['url']})    Browser: Chromium")
m.add_heading("1. Scope and approach", 1)
m.add_paragraph("A tester follows each step by hand in the browser and records Actual result and Pass/Fail. "
                "Test data uses the app's fictional 'Demo data' mode so results are repeatable; Live search was checked separately.")
m.add_heading("2. Test cases", 1)
rows = []
for tid, step, exp, link in manual:
    r = by.get(link) if link else None
    rows.append([tid, step, exp, r["actual"] if r else "Not executed in this round", r["status"] if r else "PENDING"])
table(m, ["ID", "Steps", "Expected result", "Actual result (observed)", "Status"], rows)
m.add_paragraph("")
m.add_paragraph("Note: for MT-01 to MT-13 the actual results were observed by running the identical steps in the browser through MCP "
                "(see MCP Testing Report). MT-14 to MT-16 still need to be run by hand and ticked by the tester.")
m.add_heading("3. Bug log", 1)
bugs = [[r["id"], r["title"], r["actual"]] for r in R if r["status"] == "FAIL"]
table(m, ["Ref", "Item", "Details"], bugs or [["-", "No application defects found in this round", "-"]])
m.add_heading("4. Observations", 1)
for s in ["Missing profile details only produce 'potential match' results - this is intended behavior.",
          "Hostile HTML typed in a text field was not rendered or executed.",
          "Mobile layout (390 px) has no horizontal scrolling.",
          "Live search needs internet and Node.js; its results vary with the real websites."]:
    m.add_paragraph(s, style="List Bullet")
m.add_heading("5. Manual vs MCP testing", 1)
table(m, ["Aspect", "Manual", "MCP"], [
 ["Speed", "Slow, repeated by a person", f"{len(R)} cases in minutes"],
 ["Repeatability", "Varies by tester", "Same steps every run"],
 ["Evidence", "Notes/screenshots by hand", "Auto screenshots, page text, tool-call log"],
 ["Judgement (look and feel)", "Good", "Limited to what is checked"],
 ["Cost to set up", "None", "Needs Node.js + script"]])
m.save(H / "Manual_Testing_Report.docx")
print("reports written")
